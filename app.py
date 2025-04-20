
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from docx import Document
from docx.shared import Pt
import io
import os
from datetime import datetime

app = Flask(__name__)
CORS(app)

def format_date(date_str):
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d.%m.%Y")
    except Exception:
        return date_str or ""

@app.route("/api/generate", methods=["POST"])
def generate_document():
    try:
        data = request.get_json()
        print("📥 Přijatá data:", data)

        placeholders = {
            "{{cislo_smlouvy}}": data.get("cislo_smlouvy", ""),
            "{{cislo_partnera}}": data.get("cislo_partnera", ""),
            "{{Nazev}}": data.get("Nazev", ""),
            "{{ICO}}": data.get("ICO", ""),
            "{{ulice_sidlo}}": data.get("ulice_sidlo", ""),
            "{{mesto_sidlo}}": data.get("mesto_sidlo", ""),
            "{{psc_sidlo}}": data.get("psc_sidlo", ""),
            "{{email}}": data.get("email", ""),
            "{{telefon}}": data.get("telefon", ""),
            "{{zpusob_odesilani}}": data.get("zpusob_odesilani", ""),
            "{{platby_faktury}}": data.get("platby_faktury", ""),
            "{{platby_zalohy}}": data.get("platby_zalohy", ""),
            "{{cislo_uctu}}": data.get("cislo_uctu", ""),
            "{{zahajeni_dodavek}}": format_date(data.get("zahajeni_dodavek", "")),
            "{{prolongace}}": format_date(data.get("prolongace", "")),
            "{{ean}}": data.get("ean", ""),
            "{{ulice_odber}}": data.get("ulice_odber", ""),
            "{{mesto_odber}}": data.get("mesto_odber", ""),
            "{{psc_odber}}": data.get("psc_odber", ""),
            "{{sazba}}": data.get("sazba", ""),
            "{{jistic}}": data.get("jistic", "")
        }

        doc_path = "smlouva.docx"
        if not os.path.exists(doc_path):
            return jsonify({"error": "Soubor smlouva.docx nebyl nalezen na serveru."}), 500

        template = Document(doc_path)

        for paragraph in template.paragraphs:
            for key, value in placeholders.items():
                if key in paragraph.text:
                    inline = paragraph.runs
                    for i in range(len(inline)):
                        if key in inline[i].text:
                            inline[i].text = inline[i].text.replace(key, value)
                            inline[i].font.name = 'Arial'
                            inline[i].font.size = Pt(11)

        buffer = io.BytesIO()
        template.save(buffer)
        buffer.seek(0)

        return send_file(buffer, as_attachment=True, download_name="smlouva.docx",
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/", methods=["GET"])
def home():
    return jsonify({"message": "✅ Aplikace běží. Odesílejte POST na /api/generate."})

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=8080)
